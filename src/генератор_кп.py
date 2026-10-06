import pyperclip
import tkinter as tk
from tkinter import messagebox, Checkbutton, IntVar, ttk
from datetime import datetime
import os
import random
import json
import logging

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Пути проекта и настройка логирования (общий модуль с версией для GigaChat)
from logger_config import setup_logger, SETTINGS_PATH, LEGACY_LOG_FILE

logging.getLogger("Генератор_КП").setLevel(logging.DEBUG)
logger = setup_logger("Генератор_КП", level=logging.DEBUG)

logger.info("=" * 60)
logger.info("ЗАПУСК ПРОГРАММЫ ГЕНЕРАТОР КП")
logger.info("=" * 60)

# ============================================
# ФАЙЛ ДЛЯ СОХРАНЕНИЯ НАСТРОЕК
# ============================================

SETTINGS_FILE = SETTINGS_PATH
LOG_FILE = LEGACY_LOG_FILE

def load_settings():
    """Загружает настройки из JSON файла"""
    try:
        with open(SETTINGS_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
            logger.info(f"Настройки загружены: города={len(data.get('cities', []))}, регион={data.get('region', 'не указан')}")
            return data
    except FileNotFoundError:
        logger.warning("Файл настроек не найден. Использую настройки по умолчанию.")
        return None
    except json.JSONDecodeError as e:
        logger.error(f"Ошибка чтения JSON: {e}")
        return None
    except Exception as e:
        logger.error(f"Неизвестная ошибка при загрузке настроек: {e}")
        return None

def save_settings(selected_cities, region, payment_key=None):
    """Сохраняет настройки в JSON файл"""
    if payment_key is None:
        payment_key = globals().get("PAYMENT_KEY", DEFAULT_PAYMENT_KEY)
    try:
        with open(SETTINGS_FILE, 'w', encoding='utf-8') as f:
            json.dump(
                {"cities": selected_cities, "region": region, "payment": payment_key},
                f,
                ensure_ascii=False,
                indent=2,
            )
        logger.info(f"Настройки сохранены: города={len(selected_cities)}, регион={region}, оплата={payment_key}")
        return True
    except Exception as e:
        logger.error(f"Ошибка сохранения настроек: {e}")
        return False

# ============================================
# СПИСКИ ГОРОДОВ
# ============================================

ALL_CITIES_SOUTH = [
    "Ростов-на-Дону",
    "Краснодар",
    "Новороссийск",
    "Адыгея",
    "Пятигорск",
    "Минеральные Воды"
]

ALL_CITIES_VLADIVOSTOK = [
    "Владивосток → Москва",
    "Владивосток → Новосибирск",
    "Владивосток → Краснодар",
    "Владивосток → Екатеринбург",
    "Владивосток → Казань",
    "Владивосток → Санкт-Петербург"
]

ALL_CITIES_EAST = [
    "Екатеринбург",
    "Тюмень",
    "Омск",
    "Новосибирск",
    "Красноярск",
    "Иркутск",
    "Улан-Удэ",
    "Чита",
    "Курган",
    "Кемерово",
    "Барнаул",
    "Абакан",
    "Благовещенск",
    "Владивосток",
    "Нижний Новгород",
    "Чебоксары"
]

CREATIVE_MODE = "✨ Без города (креативный режим)"

# ============================================
# ТИП ОПЛАТЫ
# ============================================
# Ключи совпадают с llm_provider.PAYMENT_TYPES, чтобы settings.json был
# одинаковым для версии с ИИ и шаблонной версии.
PAYMENT_LABELS_MAP = {
    "beznal_nds": "Безнал с НДС",
    "beznal": "Безнал без НДС",
    "cash": "Наличные",
    "discuss": "Обсуждается",
}
PAYMENT_LABELS = list(PAYMENT_LABELS_MAP.values())
DEFAULT_PAYMENT_KEY = "beznal_nds"

# Что писать про оплату для каждого варианта: фразы «Оплата … (с НДС)»
PAYMENT_PHRASES = {
    "beznal_nds": [
        "💰 Оплата по безналу (с НДС).",
        "💰 Оплата по безналу (с НДС), сроки обсуждаем.",
        "💰 Оплата — безналичный расчёт с НДС, работаем с юрлицами.",
    ],
    "beznal": [
        "💰 Оплата по безналу (без НДС).",
        "💰 Оплата по безналу (без НДС), сроки обсуждаем.",
        "💰 Безналичный расчёт без НДС, работаем с юрлицами.",
    ],
    "cash": [
        "💰 Оплата наличными (без НДС).",
        "💰 Оплата наличными (без НДС), по факту выгрузки.",
        "💰 Оплата наличными (без НДС), договор оформляем.",
    ],
    "discuss": [
        "💰 Оплата обсуждается — безнал с НДС или наличные.",
        "💰 Условия оплаты обсуждаются: безнал с НДС или наличные.",
        "💰 Оплата обсуждается — с НДС по безналу или наличными.",
    ],
}

saved = load_settings()
if saved:
    SELECTED_CITIES = saved.get("cities", ALL_CITIES_SOUTH.copy())
    CURRENT_REGION = saved.get("region", "Юга")
    PAYMENT_KEY = saved.get("payment", DEFAULT_PAYMENT_KEY)
    if PAYMENT_KEY not in PAYMENT_PHRASES:
        PAYMENT_KEY = DEFAULT_PAYMENT_KEY
    logger.info(f"Загружены настройки пользователя: регион={CURRENT_REGION}, города={SELECTED_CITIES}")
else:
    SELECTED_CITIES = ALL_CITIES_SOUTH.copy()
    CURRENT_REGION = "Юга"
    PAYMENT_KEY = DEFAULT_PAYMENT_KEY
    logger.info("Использую настройки по умолчанию: регион=Юга")

# ============================================
# ГЛОБАЛЬНЫЕ ПЕРЕМЕННЫЕ ДЛЯ GUI
# ============================================

city_vars = {}
notebook = None
text_preview = None

# ============================================
# ФУНКЦИИ
# ============================================

def get_selected_cities():
    """Возвращает список выбранных городов"""
    try:
        selected = [city for city, var in city_vars.items() if var.get() == 1]
        logger.debug(f"Выбрано городов: {len(selected)} - {selected}")
        return selected
    except Exception as e:
        logger.error(f"Ошибка получения выбранных городов: {e}")
        return []

def update_cities():
    """Обновляет выбранные города и сохраняет настройки"""
    logger.info("Обновление выбора городов")
    selected = get_selected_cities()
    region = get_current_region()
    if selected:
        global SELECTED_CITIES, CURRENT_REGION
        SELECTED_CITIES = selected
        CURRENT_REGION = region
        save_settings(selected, region)
        show_preview()
        if len(selected) == 1:
            if selected[0] == CREATIVE_MODE:
                msg = "✅ Выбран креативный режим 'Без города'"
                logger.info("Выбран креативный режим")
            else:
                msg = f"✅ Выбран город:\n\n{selected[0]}"
                logger.info(f"Выбран город: {selected[0]}")
        else:
            msg = f"✅ Выбрано: {len(selected)}\n\n{', '.join(selected)}"
            logger.info(f"Выбрано {len(selected)} городов: {', '.join(selected)}")
        messagebox.showinfo("Готово!", msg)
    else:
        logger.warning("Попытка применить выбор без выбранных городов")
        messagebox.showwarning("Внимание!", "Выберите хотя бы один город!")

def get_current_region():
    """Определяет текущий активный регион по вкладке"""
    try:
        tab = notebook.index(notebook.select())
        if tab == 0:
            return "Юга"
        elif tab == 1:
            return "Владивосток"
        else:
            return "Восток"
    except:
        return CURRENT_REGION  # Возвращаем сохранённый регион

def switch_region(event):
    """Переключает регион при смене вкладки"""
    region = get_current_region()
    global CURRENT_REGION, SELECTED_CITIES
    CURRENT_REGION = region
    logger.info(f"Переключение на регион: {region}")
    
    if region == "Владивосток":
        cities = ALL_CITIES_VLADIVOSTOK
    elif region == "Восток":
        cities = ALL_CITIES_EAST + [CREATIVE_MODE]
    else:
        cities = ALL_CITIES_SOUTH
    
    # Обновляем чекбоксы
    for city, var in city_vars.items():
        var.set(1 if city in cities else 0)
    
    selected = get_selected_cities()
    if selected:
        SELECTED_CITIES = selected
        save_settings(selected, region)
    show_preview()

# ============================================
# ОФОРМЛЕНИЕ ГОРОДОВ - РАСШИРЕННЫЙ КРЕАТИВ
# ============================================

def format_route():
    """Форматирует маршрут с креативным оформлением"""
    cities = get_selected_cities()
    region = get_current_region()
    
    is_creative = CREATIVE_MODE in cities and len(cities) == 1
    
    if is_creative and region == "Восток":
        creative_phrases = [
            "🚛 Лот машин на автовоз: Москва → Восточные регионы!",
            "🚛 Лот машин ждёт автовоз: Москва → Восток!",
            "🚛 Автовозам: лот машин Москва → Восточная Сибирь и ДВ!",
            "🚛 Лот машин на автовоз: Москва → Урал и Сибирь!",
            "🚛 Груз для автовоза: Москва → Урал, Сибирь, Дальний Восток!",
            "🚛 Лот машин на автовоз: Москва → Восточные города России!",
            "🚛 Лот машин ждёт автовоз: от Урала до Тихого океана!",
            "🚛 Лот машин на автовоз: Москва → Восточные регионы РФ!",
            "🚛 Груз для автовоза: Москва → Дальний Восток!",
            "🚛 Лот машин на автовоз — восточный маршрут!",
        ]
        return random.choice(creative_phrases)
    
    if not cities:
        return "Нет выбранных городов"
    
    cities = [c for c in cities if c != CREATIVE_MODE]
    if not cities:
        if region == "Восток":
            creative_phrases = [
                "🚛 Лот машин на автовоз: Москва → Восточные регионы!",
                "🚛 Лот машин ждёт автовоз: Москва → Восток!",
                "🚛 Автовозам: лот машин Москва → Восточная Сибирь и ДВ!",
                "🚛 Лот машин на автовоз: Москва → Урал и Сибирь!",
                "🚛 Груз для автовоза: Москва → Урал, Сибирь, Дальний Восток!",
                "🚛 Лот машин на автовоз: Москва → Восточные города России!",
                "🚛 Лот машин ждёт автовоз: от Урала до Тихого океана!",
                "🚛 Лот машин на автовоз: Москва → Восточные регионы РФ!",
            ]
            return random.choice(creative_phrases)
        return "Нет выбранных городов"
    
    if len(cities) == 1:
        city = cities[0]
        one_city_styles = [
            f"🚛 Лот машин на автовоз: Москва → {city}!",
            f"🚛 Лот машин ждёт автовоз: Москва → {city}!",
            f"🚛 Выгрузка: {city} — ищем автовоз на этот лот!",
            f"🚛 Автовозам: лот машин Москва → {city}!",
            f"🚛 Машины ждут автовоз: Москва → {city}!",
            f"🚛 Груз для автовоза: Москва → {city}!",
            f"🚛 Лот сформирован, ищем автовоз: Москва → {city}!",
            f"🚛 Есть машины на перевозку автовозом: Москва → {city}!",
            f"🚛 Лот машин под автовоз: Москва → {city}!",
            f"🚛 Машины на отправку: Москва → {city}, нужен автовоз!",
        ]
        return random.choice(one_city_styles)
    
    styles = [
        lambda: "\n".join([f"{i+1}. {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"{i+1}) {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"• {city}" for city in cities]),
        lambda: "\n".join([f"▸ {city}" for city in cities]),
        lambda: "\n".join([f"■ {city}" for city in cities]),
        lambda: "\n".join([f"✦ {city}" for city in cities]),
        lambda: "\n".join([f"{chr(1072 + i)}) {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"{chr(1040 + i)}) {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"{['I','II','III','IV','V','VI','VII','VIII','IX','X','XI','XII','XIII','XIV','XV','XVI'][i]}. {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"➜ {city}" for city in cities]),
        lambda: "\n".join([f"✅ {city}" for city in cities]),
        lambda: "\n".join([f"★ {city}" for city in cities]),
        lambda: ", ".join(cities),
        lambda: "\n".join([f">> {city}" for city in cities]),
        lambda: "\n".join([f"[{i+1}] {city}" for i, city in enumerate(cities)]),
        lambda: "\n".join([f"{i+1}.     {city}" for i, city in enumerate(cities)]),
    ]
    return random.choice(styles)()

# ============================================
# КРЕАТИВНЫЕ ФРАЗЫ - РАСШИРЕННЫЕ
# ============================================
#
# Специфика проекта: машины в лоте — это ГРУЗ для автовоза, а не наш транспорт.
# Поэтому в каждой фразе про машины стоит контекст «на автовоз»: без него текст
# читается как «наши авто готовы к отправке» (см. README, раздел про автовозы).

INTROS_SOUTH = [
    "🚛 Есть лот машин на автовоз: Москва → Юга!",
    "🚛 Свежий лот на автовоз — Москва → Юг России!",
    "🚛 Лот машин ждёт автовоз — Москва → Юга!",
    "🚛 Автовозам: свежий лот с Москвы на Юга!",
    "🚛 Машины на отправку — нужен автовоз, Москва → Юга!",
    "🚛 Лот сформирован, ищем автовоз: Москва → Юга!",
    "🚛 Груз для автовоза — Москва → южные города!",
    "🚛 Есть машины на перевозку автовозом: Москва → Юга!",
    "🚛 Новый лот машин под автовоз — Москва → Юга!",
    "🚛 Автовозам на заметку: лот машин с Москвы на Юга!",
    "🚛 Лот машин на юг — загрузка на автовоз!",
    "🚛 Машины ждут автовоз: Москва → Юга!",
]

CARGO_SOUTH = [
    "🚗 Машины готовы к погрузке на автовоз.",
    "🚗 Лот сформирован, машины ждут автовоз.",
    "🚗 Автомобили для перевозки — ищем автовоз.",
    "🚗 Есть машины на перевозку автовозом.",
    "🚗 Машины ждут загрузки на автовоз.",
    "🚗 Лот машин ждёт автовоз — загрузка быстрая.",
    "🚗 Машины на площадке, нужен автовоз под погрузку.",
    "🚗 Машины ждут своей очереди на автовоз.",
    "🚗 Лот машин сформирован — ищем автовоз.",
]

PAYMENT_SOUTH = [
    "💰 Оплата по безналу (с НДС), сроки обсуждаем.",
    "💰 Оплата по безналу (с НДС), возможна предоплата.",
    "💰 Безналичный расчёт с НДС, гибкие условия.",
    "💰 Оплата по безналу (с НДС), работаем с юрлицами.",
    "💰 Оплата по безналу (с НДС), договор сразу.",
    "💰 Оплата по безналу (с НДС), закрывающие документы.",
    "💰 Оплата по безналу (с НДС), сроки подберём.",
]

LEGAL_SOUTH = [
    "📋 Работаем по договору, платим без задержек.",
    "📄 Заключаем договор, выплаты строго в срок, без просрочек.",
    "🤝 Официальное оформление, оплата вовремя, гарантируем надёжность.",
    "📑 Работаем официально, платежи без задержек, договор гарантирует защиту.",
    "📋 Все сделки по договору, выплаты чётко в срок, без задержек.",
    "📄 Юридически чистая работа — договор, оплата по факту, без просрочек.",
    "⚖️ Полное юридическое сопровождение — ваша безопасность в приоритете!",
]

CTA_SOUTH = [
    "✍️ Все детали и ставки — в личку. Пишите! 📩",
    "📩 Ждём ваших сообщений для обсуждения деталей и ставок!",
    "💬 Пишите в личные сообщения — обсудим все условия и ставки!",
    "📲 Свяжитесь с нами в личных сообщениях — пришлём ставки!",
    "✍️ Детали и стоимость — в личку. Ждём ваших вопросов!",
    "📩 Обсудим все нюансы и ставки в личных сообщениях!",
    "📱 Ждём ваших сообщений — готовы ответить на все вопросы!",
]

INTROS_VLADIVOSTOK = [
    "🚛 Лот машин на автовоз: Владивосток → по России!",
    "🚛 Свежий лот на автовоз — из Владивостока по России!",
    "🚛 Автовозам: лот машин с Владивостока в любой город!",
    "🚛 Лот машин ждёт автовоз — Владивосток → по России!",
    "🚛 Машины на отправку с Владивостока — нужен автовоз!",
    "🚛 Лот сформирован, ищем автовоз: Владивосток → Россия!",
    "🚛 Груз для автовоза — из Владивостока по всей стране!",
    "🚛 Есть машины на перевозку автовозом: Владивосток → Россия!",
    "🚛 Новый лот машин под автовоз — старт во Владивостоке!",
    "🚛 Автовозам на заметку: лот машин с Дальнего Востока!",
    "🚛 Лот машин с Владивостока — загрузка на автовоз!",
    "🚛 Машины ждут автовоз: Владивосток → любой регион!",
]

CARGO_VLADIVOSTOK = [
    "🚗 Машины готовы к погрузке на автовоз.",
    "🚗 Лот сформирован, машины ждут автовоз.",
    "🚗 Автомобили для перевозки — ищем автовоз.",
    "🚗 Есть машины на перевозку автовозом.",
    "🚗 Машины ждут загрузки на автовоз.",
    "🚗 Лот машин ждёт автовоз — погрузка быстрая.",
]

PAYMENT_VLADIVOSTOK = [
    "💰 Оплата наличными (без НДС) — 50/50 при погрузке.",
    "💰 Оплата наличными (без НДС) — 50% при загрузке, 50% по факту.",
    "💰 Оплата по безналу (с НДС) — возможна предоплата.",
    "💰 Оплата по безналу (с НДС), работаем с юрлицами.",
    "💰 Условия оплаты обсуждаются — наличные или безнал с НДС.",
]

LEGAL_VLADIVOSTOK = [
    "📦 Работаем чисто, без задержек. Подробности по каждому лоту — в личные сообщения.",
    "📋 Работаем по договору, все сделки прозрачны.",
    "📄 Гарантируем чистоту сделок, работаем без задержек.",
    "🤝 Надёжные партнёры — оплата вовремя, без задержек.",
    "📑 Полная юридическая прозрачность — ваше спокойствие гарантируем!",
]

CTA_VLADIVOSTOK = [
    "✍️ Пишите — обсудим ставки, даты и условия!",
    "📩 Обсудим детали в личных сообщениях — ждём ваших вопросов!",
    "💬 Свяжитесь с нами для уточнения ставок и дат!",
    "📲 Пишите — договоримся о всех условиях!",
    "📱 Ждём вашего сообщения — готовы к диалогу!",
]

INTROS_EAST = [
    "🚛 Лот машин на автовоз: Москва → Восток!",
    "🚛 Свежий лот на автовоз — Москва → Урал и Сибирь!",
    "🚛 Автовозам: лот машин на восточное направление!",
    "🚛 Лот машин ждёт автовоз — Москва → Восток!",
    "🚛 Машины на отправку на Восток — нужен автовоз!",
    "🚛 Лот сформирован, ищем автовоз: Москва → Восток!",
    "🚛 Груз для автовоза — Москва → Урал → Сибирь → ДВ!",
    "🚛 Есть машины на перевозку автовозом на Восток!",
    "🚛 Новый лот машин под автовоз — восточный маршрут!",
    "🚛 Автовозам на заметку: лот машин с Москвы на Восток!",
    "🚛 Лот машин на восточные регионы — загрузка на автовоз!",
    "🚛 Машины ждут автовоз: Москва → Восток!",
]

CARGO_EAST = [
    "🚗 Машины готовы к погрузке на автовоз.",
    "🚗 Лот сформирован, машины ждут автовоз.",
    "🚗 Автомобили для перевозки — ищем автовоз.",
    "🚗 Есть машины на перевозку автовозом.",
    "🚗 Машины ждут загрузки на автовоз.",
]

PAYMENT_EAST = [
    "💰 Оплата по безналу (с НДС), сроки обсуждаем.",
    "💰 Оплата по безналу (с НДС), гибкие условия.",
    "💰 Безналичный расчёт с НДС — быстро и прозрачно.",
    "💰 Оплата по безналу (с НДС), условия под каждого.",
    "💰 Условия оплаты обсуждаются — наличные или безнал с НДС.",
]

LEGAL_EAST = [
    "📋 Работаем по договору, выплаты в срок.",
    "📄 Заключаем договор — гарантия надёжности.",
    "🤝 Официальное оформление, оплата вовремя.",
    "📑 Все сделки по договору, без задержек.",
    "⚖️ Полное юридическое сопровождение сделки!",
]

CTA_EAST = [
    "✍️ Детали и ставки — в ЛС. Ждём ваших сообщений!",
    "📩 Обсудим условия и ставки в личных сообщениях.",
    "💬 Пишите — рассчитаем оптимальные условия!",
    "📲 Свяжитесь с нами в личных сообщениях — уточним детали и ставки!",
    "📱 Ждём вашего звонка или сообщения!",
]

# ============================================
# ФУНКЦИЯ ГЕНЕРАЦИИ
# ============================================

def payment_phrase():
    """Фраза про оплату по выбранному в интерфейсе типу (с НДС / без НДС / наличные / обсуждается)."""
    phrases = PAYMENT_PHRASES.get(PAYMENT_KEY, PAYMENT_PHRASES[DEFAULT_PAYMENT_KEY])
    return random.choice(phrases)


def get_payment_key():
    """Читает тип оплаты из выпадающего списка. До создания виджета — сохранённое значение."""
    try:
        label = payment_var.get()
    except NameError:
        return PAYMENT_KEY
    for key, shown in PAYMENT_LABELS_MAP.items():
        if shown == label:
            return key
    return PAYMENT_KEY


def update_payment(_event=None):
    """Обработчик выбора оплаты: сохраняет настройку и обновляет предпросмотр."""
    global PAYMENT_KEY
    PAYMENT_KEY = get_payment_key()
    logger.info(f"Выбран тип оплаты: {PAYMENT_KEY}")
    save_settings(get_selected_cities(), get_current_region(), PAYMENT_KEY)
    show_preview()


def generate_creative_text():
    """Генерирует креативный текст коммерческого предложения"""
    region = get_current_region()
    route = format_route()
    logger.debug(f"Генерация текста для региона: {region}")
    
    if region == "Владивосток":
        intro = random.choice(INTROS_VLADIVOSTOK)
        cargo = random.choice(CARGO_VLADIVOSTOK)
        legal = random.choice(LEGAL_VLADIVOSTOK)
        cta = random.choice(CTA_VLADIVOSTOK)
    elif region == "Восток":
        intro = random.choice(INTROS_EAST)
        cargo = random.choice(CARGO_EAST)
        legal = random.choice(LEGAL_EAST)
        cta = random.choice(CTA_EAST)
    else:
        intro = random.choice(INTROS_SOUTH)
        cargo = random.choice(CARGO_SOUTH)
        legal = random.choice(LEGAL_SOUTH)
        cta = random.choice(CTA_SOUTH)

    payment = payment_phrase()

    result = f"{intro}\n\n{route}\n\n{cargo}\n{payment}\n{legal}\n\n{cta}"
    logger.debug("Текст успешно сгенерирован")
    return result

# ============================================
# КНОПКИ
# ============================================

def copy_to_clipboard():
    """Копирует сгенерированный текст в буфер обмена"""
    logger.info("Копирование текста в буфер обмена")
    try:
        text = generate_creative_text()
        pyperclip.copy(text)
        logger.debug("Текст скопирован в буфер обмена")
        messagebox.showinfo("Готово!", "✅ Текст скопирован в буфер обмена!")
    except Exception as e:
        logger.error(f"Ошибка копирования в буфер: {e}")
        messagebox.showerror("Ошибка", f"❌ {str(e)}")

def save_to_word():
    """Сохраняет сгенерированный текст в Word документ"""
    logger.info("Сохранение в Word")
    try:
        text = generate_creative_text()
        doc = Document()
        title = doc.add_heading('КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ', level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_paragraph = doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y')}")
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()
        for line in text.split('\n'):
            if line.strip():
                doc.add_paragraph(line)
        region = get_current_region()
        filename = f"КП_{region}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(os.path.expanduser('~'), 'Desktop', filename)
        doc.save(filepath)
        logger.info(f"Word-файл сохранён: {filepath}")
        messagebox.showinfo("Готово!", f"✅ Word-файл создан на рабочем столе!\n\n{filename}")
    except Exception as e:
        logger.error(f"Ошибка сохранения в Word: {e}")
        messagebox.showerror("Ошибка", f"❌ {str(e)}")

def save_multiple_variants():
    """Создаёт 3 разных варианта в Word"""
    logger.info("Создание 3 вариантов в Word")
    try:
        region = get_current_region()
        for i in range(3):
            text = generate_creative_text()
            doc = Document()
            title = doc.add_heading('КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ', level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            date_paragraph = doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y')}")
            date_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph()
            for line in text.split('\n'):
                if line.strip():
                    doc.add_paragraph(line)
            filename = f"КП_{region}_вариант_{i+1}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            filepath = os.path.join(os.path.expanduser('~'), 'Desktop', filename)
            doc.save(filepath)
            logger.debug(f"Вариант {i+1} сохранён: {filepath}")
        messagebox.showinfo("Готово!", "✅ 3 файла созданы на рабочем столе!")
    except Exception as e:
        logger.error(f"Ошибка создания 3 вариантов: {e}")
        messagebox.showerror("Ошибка", f"❌ {str(e)}")

def show_preview():
    """Обновляет предпросмотр текста"""
    logger.debug("Обновление предпросмотра")
    try:
        text = generate_creative_text()
        text_preview.config(state="normal")
        text_preview.delete("1.0", tk.END)
        text_preview.insert("1.0", text)
        text_preview.config(state="disabled")
    except Exception as e:
        logger.error(f"Ошибка обновления предпросмотра: {e}")

def create_installer():
    """Создаёт установщик с помощью Inno Setup"""
    logger.info("Запуск создания установщика")
    try:
        import subprocess
        import shutil
        current_dir = os.path.dirname(os.path.abspath(__file__))
        desktop = os.path.join(os.path.expanduser('~'), 'Desktop')
        messagebox.showinfo("Установщик", "🔨 Сборка...\n\nПодождите 2-3 минуты.")
        subprocess.run(['pyinstaller', '--onefile', '--console', '--name', 'Генератор_КП', os.path.join(current_dir, 'генератор_кп.py')], check=True)
        exe_source = os.path.join(current_dir, 'dist', 'Генератор_КП.exe')
        setup_folder = os.path.join(desktop, 'Setup')
        if not os.path.exists(setup_folder):
            os.makedirs(setup_folder)
        shutil.copy(exe_source, os.path.join(setup_folder, 'Генератор_КП.exe'))
        iss_content = '''[Setup]
AppName=Генератор КП
AppVersion=1.0
DefaultDirName={pf}\\Генератор КП
DefaultGroupName=Генератор КП
OutputDir=.
OutputBaseFilename=Генератор_КП_Установщик
Compression=lzma
SolidCompression=yes
WizardStyle=modern

[Files]
Source: "Генератор_КП.exe"; DestDir: "{app}"

[Icons]
Name: "{group}\\Генератор КП"; Filename: "{app}\\Генератор_КП.exe"
Name: "{commondesktop}\\Генератор КП"; Filename: "{app}\\Генератор_КП.exe"

[Run]
Filename: "{app}\\Генератор_КП.exe"; Description: "Запустить Генератор КП"; Flags: postinstall nowait skipifsilent'''
        with open(os.path.join(setup_folder, 'setup.iss'), 'w', encoding='utf-8') as f:
            f.write(iss_content)
        os.startfile(setup_folder)
        logger.info("Папка Setup создана на рабочем столе")
        messagebox.showinfo("Готово!", "✅ Папка Setup создана на рабочем столе!\n\nОткройте setup.iss в Inno Setup и нажмите F9")
    except Exception as e:
        logger.error(f"Ошибка создания установщика: {e}")
        messagebox.showerror("Ошибка", f"❌ {str(e)}\n\nУстановите PyInstaller:\npip install pyinstaller")

def on_closing():
    """Обработчик закрытия окна - сохраняет настройки"""
    logger.info("Закрытие программы - сохранение настроек")
    try:
        # Сохраняем текущие настройки перед выходом
        selected = get_selected_cities()
        region = get_current_region()
        
        logger.info(f"Сохраняю перед выходом: регион={region}, городов={len(selected)}")
        
        if selected:
            success = save_settings(selected, region)
            if success:
                logger.info("Настройки успешно сохранены перед выходом")
            else:
                logger.error("Не удалось сохранить настройки перед выходом")
        else:
            # Если ничего не выбрано, сохраняем последние выбранные
            if SELECTED_CITIES:
                logger.info(f"Сохраняю последние выбранные города: {SELECTED_CITIES}")
                save_settings(SELECTED_CITIES, CURRENT_REGION)
            else:
                logger.warning("Нет городов для сохранения")
    except Exception as e:
        logger.error(f"Ошибка сохранения настроек при выходе: {e}")
    
    window.destroy()
    logger.info("Программа завершена")

# ============================================
# ГРАФИЧЕСКИЙ ИНТЕРФЕЙС
# ============================================

logger.info("Создание графического интерфейса")

window = tk.Tk()
window.title("Генератор КП - Автоперевозки")
window.geometry("750x850")
window.resizable(False, False)
BG_COLOR = "#f0f4f8"
window.configure(bg=BG_COLOR)

# Обработчик закрытия окна
window.protocol("WM_DELETE_WINDOW", on_closing)

tk.Label(window, text="🚛 ГЕНЕРАТОР КОММЕРЧЕСКИХ ПРЕДЛОЖЕНИЙ", font=("Arial", 16, "bold"), bg=BG_COLOR, pady=10).pack()

notebook = ttk.Notebook(window)
notebook.pack(fill="x", padx=20, pady=5)

tab_south = tk.Frame(notebook, bg=BG_COLOR)
notebook.add(tab_south, text="📍 Юга")

tab_vladivostok = tk.Frame(notebook, bg=BG_COLOR)
notebook.add(tab_vladivostok, text="📍 Владивосток")

tab_east = tk.Frame(notebook, bg=BG_COLOR)
notebook.add(tab_east, text="📍 Восток")

notebook.bind("<<NotebookTabChanged>>", switch_region)

city_vars = {}

# ===== ЮГА =====
tk.Label(tab_south, text="📍 Выберите города для маршрута:", font=("Arial", 10, "bold"), bg=BG_COLOR).grid(row=0, column=0, columnspan=3, pady=5)
row = 1
for i, city in enumerate(ALL_CITIES_SOUTH):
    var = IntVar(value=1 if city in SELECTED_CITIES and CURRENT_REGION == "Юга" else 0)
    city_vars[city] = var
    Checkbutton(tab_south, text=city, variable=var, font=("Arial", 10), bg=BG_COLOR, cursor="hand2").grid(row=row, column=i % 3, padx=15, pady=3, sticky="w")
    if i % 3 == 2:
        row += 1

# ===== ВЛАДИВОСТОК =====
tk.Label(tab_vladivostok, text="📍 Выберите направления:", font=("Arial", 10, "bold"), bg=BG_COLOR).grid(row=0, column=0, columnspan=3, pady=5)
row_vl = 1
for i, city in enumerate(ALL_CITIES_VLADIVOSTOK):
    var = IntVar(value=1 if city in SELECTED_CITIES and CURRENT_REGION == "Владивосток" else 0)
    city_vars[city] = var
    Checkbutton(tab_vladivostok, text=city, variable=var, font=("Arial", 10), bg=BG_COLOR, cursor="hand2").grid(row=row_vl, column=i % 3, padx=15, pady=3, sticky="w")
    if i % 3 == 2:
        row_vl += 1

# ===== ВОСТОК =====
tk.Label(tab_east, text="📍 Выберите города для маршрута (в порядке следования):", font=("Arial", 10, "bold"), bg=BG_COLOR).grid(row=0, column=0, columnspan=4, pady=5)

var_creative = IntVar(value=1 if CREATIVE_MODE in SELECTED_CITIES and CURRENT_REGION == "Восток" else 0)
city_vars[CREATIVE_MODE] = var_creative
Checkbutton(tab_east, text=CREATIVE_MODE, variable=var_creative, font=("Arial", 9, "italic"), bg=BG_COLOR, fg="#FF5722", cursor="hand2").grid(row=1, column=0, columnspan=4, padx=10, pady=5, sticky="w")

row_east = 2
for i, city in enumerate(ALL_CITIES_EAST):
    var = IntVar(value=1 if city in SELECTED_CITIES and CURRENT_REGION == "Восток" else 0)
    city_vars[city] = var
    Checkbutton(tab_east, text=city, variable=var, font=("Arial", 9), bg=BG_COLOR, cursor="hand2").grid(row=row_east, column=i % 4, padx=10, pady=2, sticky="w")
    if i % 4 == 3:
        row_east += 1

# Устанавливаем правильную вкладку при запуске
if CURRENT_REGION == "Юга":
    notebook.select(0)
elif CURRENT_REGION == "Владивосток":
    notebook.select(1)
else:
    notebook.select(2)

payment_frame = tk.Frame(window, bg=BG_COLOR)
payment_frame.pack(pady=3)

tk.Label(payment_frame, text="💰 Оплата:", font=("Arial", 10, "bold"), bg=BG_COLOR).pack(side="left", padx=5)

payment_var = tk.StringVar(value=PAYMENT_LABELS_MAP.get(PAYMENT_KEY, PAYMENT_LABELS_MAP[DEFAULT_PAYMENT_KEY]))
payment_combo = ttk.Combobox(payment_frame, textvariable=payment_var, values=PAYMENT_LABELS, state="readonly", width=18, font=("Arial", 9))
payment_combo.pack(side="left", padx=5)
payment_combo.bind("<<ComboboxSelected>>", update_payment)

tk.Button(window, text="✅ ПРИМЕНИТЬ ВЫБОР ГОРОДОВ", font=("Arial", 9, "bold"), bg="#4CAF50", fg="white", padx=20, pady=5, cursor="hand2", command=update_cities).pack(pady=5)
tk.Frame(window, height=2, bg="#ccc").pack(fill="x", padx=20, pady=5)

tk.Button(window, text="🔄 ПОКАЗАТЬ НОВЫЙ ВАРИАНТ", font=("Arial", 10, "bold"), bg="#9C27B0", fg="white", padx=20, pady=8, cursor="hand2", command=show_preview).pack(pady=5)

text_preview = tk.Text(window, height=12, font=("Arial", 10), wrap="word", relief="solid", bd=1, bg="white")
text_preview.pack(fill="both", padx=20, pady=10, expand=True)
show_preview()

tk.Label(window, text="📌 Для Востока: отметьте 'Без города' для креативного режима без перечисления городов", font=("Arial", 9), bg=BG_COLOR, fg="#888").pack(pady=2)

button_frame = tk.Frame(window, bg=BG_COLOR)
button_frame.pack(pady=5)

tk.Button(button_frame, text="📋 КОПИРОВАТЬ В БУФЕР", font=("Arial", 11, "bold"), bg="#4CAF50", fg="white", padx=20, pady=10, width=20, cursor="hand2", command=copy_to_clipboard).grid(row=0, column=0, padx=5, pady=5)
tk.Button(button_frame, text="📄 СОХРАНИТЬ В WORD", font=("Arial", 11, "bold"), bg="#2196F3", fg="white", padx=20, pady=10, width=20, cursor="hand2", command=save_to_word).grid(row=0, column=1, padx=5, pady=5)

tk.Button(window, text="📚 СОЗДАТЬ 3 РАЗНЫХ ВАРИАНТА В WORD", font=("Arial", 10, "bold"), bg="#FF9800", fg="white", padx=20, pady=10, width=45, cursor="hand2", command=save_multiple_variants).pack(pady=5)

tk.Button(window, text="🛠 СОЗДАТЬ УСТАНОВЩИК (Inno Setup)", font=("Arial", 10, "bold"), bg="#FF5722", fg="white", padx=20, pady=10, width=45, cursor="hand2", command=create_installer).pack(pady=5)

tk.Button(window, text="✖ ВЫХОД", font=("Arial", 10), bg="#f44336", fg="white", padx=15, pady=5, width=15, cursor="hand2", command=on_closing).pack(pady=10)

logger.info("Графический интерфейс загружен. Запуск mainloop()")
window.mainloop()